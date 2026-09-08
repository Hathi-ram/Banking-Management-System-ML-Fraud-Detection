# ==========================================================
# IMPORTS
# ==========================================================

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    session,
    flash,
    send_file
)

from functools import wraps
from models.database import get_connection

from openpyxl import Workbook
from io import BytesIO

import tempfile
import os
import subprocess

from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors


# ==========================================================
# IMPORT BLUEPRINTS
# ==========================================================

from routes.customer import customer_bp
from routes.account import account_bp
from routes.transaction import transaction_bp
from routes.loan import loan_bp
from routes.card import card_bp
from routes.employee import employee_bp
from routes.activity import activity_bp
from routes.admin import admin_bp
from routes.report import report_bp


# ==========================================================
# FLASK APP
# ==========================================================

app = Flask(__name__)

app.secret_key = "banking_management_system_secret_key_2026"


# ==========================================================
# REGISTER BLUEPRINTS
# ==========================================================

app.register_blueprint(customer_bp)
app.register_blueprint(account_bp)
app.register_blueprint(transaction_bp)
app.register_blueprint(loan_bp)
app.register_blueprint(card_bp)
app.register_blueprint(employee_bp)
app.register_blueprint(activity_bp)
app.register_blueprint(admin_bp)
app.register_blueprint(report_bp)


# ==========================================================
# PROTECT BANKING PAGES
# ==========================================================

@app.before_request
def protect_pages():

    public_routes = [
        "home",
        "login",
        "static"
    ]

    # Allow public pages
    if request.endpoint in public_routes:
        return

    # Require login for all other pages
    if "admin_id" not in session:
        return redirect("/login")


# ==========================================================
# HOME
# ==========================================================

@app.route("/")
def home():

    connection = get_connection()

    if connection:

        db_name = connection.database

        connection.close()

        return render_template(
            "home.html",
            database=db_name
        )

    return "Database Connection Failed"


# ==========================================================
# LOGIN
# ==========================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        username = request.form["username"]
        password = request.form["password"]

        connection = get_connection()

        cursor = connection.cursor(dictionary=True)

        cursor.execute("""
            SELECT *
            FROM Admin
            WHERE username=%s
            AND password=%s
        """, (username, password))

        admin = cursor.fetchone()

        cursor.close()
        connection.close()

        if admin:

            session["admin_id"] = admin["admin_id"]
            session["username"] = admin["username"]

            log_activity(
                "Admin '" +
                admin["username"] +
                "' logged into the system"
            )

            return redirect("/dashboard")

        else:

            return render_template(
                "login.html",
                error="Invalid username or password"
            )

    return render_template("login.html")


# ==========================================================
# ACTIVITY LOGGER
# ==========================================================

def log_activity(message):

    connection = None
    cursor = None

    try:

        connection = get_connection()

        cursor = connection.cursor()

        cursor.execute(
            """
            INSERT INTO Activity_Log
            (
                activity
            )
            VALUES
            (
                %s
            )
            """,
            (message,)
        )

        connection.commit()

    except Exception as e:

        print("Activity Log Error:", e)

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


# ==========================================================
# LOGOUT
# ==========================================================

@app.route("/logout")
def logout():

    # Get logged-in admin username
    username = session.get("username")

    connection = get_connection()

    if connection:

        cursor = connection.cursor()

        cursor.execute(
            """
            INSERT INTO Activity_Log
            (
                activity
            )
            VALUES
            (
                %s
            )
            """,
            (
                f"Admin '{username}' logged out of the system",
            )
        )

        connection.commit()

        cursor.close()
        connection.close()

    # Clear login session
    session.clear()

    return redirect("/login")


# ==========================================================
# LOGIN REQUIRED DECORATOR
# ==========================================================

def login_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        if "admin_id" not in session:
            return redirect("/login")

        return function(*args, **kwargs)

    return wrapper


# ==========================================================
# DASHBOARD
# ==========================================================

@app.route("/dashboard")
def dashboard():

    if "admin_id" not in session:
        return redirect("/login")

    connection = None
    cursor = None

    try:

        connection = get_connection()

        cursor = connection.cursor()

        # ======================================================
        # BANKING OVERVIEW
        # ======================================================

        # -------------------------
        # Total Customers
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Customers
        """)

        customers = cursor.fetchone()[0]


        # -------------------------
        # Total Accounts
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Accounts
        """)

        accounts = cursor.fetchone()[0]


        # -------------------------
        # Total Transactions
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
        """)

        transactions = cursor.fetchone()[0]


        # -------------------------
        # Total Loans
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Loans
        """)

        loans = cursor.fetchone()[0]


        # -------------------------
        # Total Employees
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Employees
        """)

        employees = cursor.fetchone()[0]


        # -------------------------
        # Total Cards
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Cards
        """)

        cards = cursor.fetchone()[0]


        # -------------------------
        # Active Cards
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Cards
            WHERE card_status = 'Active'
        """)

        active_cards = cursor.fetchone()[0]


        # -------------------------
        # Total Bank Balance
        # -------------------------

        cursor.execute("""
            SELECT IFNULL(SUM(balance), 0)
            FROM Accounts
        """)

        total_balance = cursor.fetchone()[0]


        # -------------------------
        # Total Approved Loan Amount
        # -------------------------

        cursor.execute("""
            SELECT IFNULL(SUM(loan_amount), 0)
            FROM Loans
            WHERE loan_status = 'Approved'
        """)

        total_loans = cursor.fetchone()[0]


        # ======================================================
        # RECENT TRANSACTIONS
        # ======================================================

        cursor.execute("""
            SELECT
                transaction_reference,
                transaction_type,
                amount,
                transaction_date,
                ml_prediction,
                ml_risk_score,
                ml_risk_level
            FROM Transactions
            ORDER BY transaction_date DESC
            LIMIT 5
        """)

        recent_transactions = cursor.fetchall()


        # ======================================================
        # ML FRAUD DETECTION ANALYTICS
        # ======================================================

        # -------------------------
        # Transactions Analyzed
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
        """)

        ml_analyzed = cursor.fetchone()[0]


        # -------------------------
        # Normal Transactions
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_prediction = 0
        """)

        ml_normal = cursor.fetchone()[0]


        # -------------------------
        # Suspicious Transactions
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_prediction = 1
        """)

        ml_suspicious = cursor.fetchone()[0]


        # -------------------------
        # High Risk Transactions
        # -------------------------

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_risk_level = 'High'
        """)

        ml_high_risk = cursor.fetchone()[0]


        # ======================================================
        # HIGH-RISK TRANSACTIONS
        # ======================================================

        cursor.execute("""
            SELECT
                transaction_reference,
                transaction_type,
                amount,
                transaction_date,
                ml_prediction,
                ml_risk_score,
                ml_risk_level
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_risk_level = 'High'
            ORDER BY ml_risk_score DESC
            LIMIT 5
        """)

        high_risk_transactions = cursor.fetchall()


        # ======================================================
        # TOP 5 CUSTOMERS
        # ======================================================

        cursor.execute("""
            SELECT
                c.first_name,
                c.last_name,
                a.account_number,
                a.balance
            FROM Customers c
            INNER JOIN Accounts a
                ON c.customer_id = a.customer_id
            ORDER BY a.balance DESC
            LIMIT 5
        """)

        top_customers = cursor.fetchall()


        # ======================================================
        # BRANCH PERFORMANCE
        # ======================================================

        cursor.execute("""
            SELECT
                b.branch_name,
                COUNT(DISTINCT a.customer_id) AS customers,
                COUNT(DISTINCT a.account_number) AS accounts,
                COUNT(DISTINCT l.loan_id) AS loans
            FROM Branches b

            LEFT JOIN Accounts a
                ON b.branch_id = a.branch_id

            LEFT JOIN Loans l
                ON a.customer_id = l.customer_id

            GROUP BY
                b.branch_id,
                b.branch_name

            ORDER BY
                b.branch_name
        """)

        branch_performance = cursor.fetchall()


        # ======================================================
        # AVERAGE ACCOUNT BALANCE
        # ======================================================

        cursor.execute("""
            SELECT IFNULL(
                ROUND(AVG(balance), 2),
                0
            )
            FROM Accounts
        """)

        avg_balance = cursor.fetchone()[0]


        # ======================================================
        # ACTIVE ACCOUNTS
        # ======================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Accounts
            WHERE account_status = 'Active'
        """)

        active_accounts = cursor.fetchone()[0]


        # ======================================================
        # TOTAL BRANCHES
        # ======================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Branches
        """)

        branches = cursor.fetchone()[0]


    except Exception as e:

        print("Dashboard Error:", e)

        flash(
            "Unable to load dashboard data.",
            "danger"
        )

        customers = 0
        accounts = 0
        transactions = 0
        loans = 0
        employees = 0
        cards = 0
        active_cards = 0
        total_balance = 0
        total_loans = 0

        recent_transactions = []

        ml_analyzed = 0
        ml_normal = 0
        ml_suspicious = 0
        ml_high_risk = 0

        high_risk_transactions = []

        top_customers = []
        branch_performance = []

        avg_balance = 0
        active_accounts = 0
        branches = 0


    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()


    # ======================================================
    # RENDER DASHBOARD
    # ======================================================

    return render_template(
        "dashboard.html",

        # Banking Overview
        customers=customers,
        accounts=accounts,
        transactions=transactions,
        loans=loans,
        employees=employees,
        cards=cards,
        active_cards=active_cards,
        total_balance=total_balance,
        total_loans=total_loans,

        # Recent Transactions
        recent_transactions=recent_transactions,

        # Customers
        top_customers=top_customers,

        # Branches
        branch_performance=branch_performance,

        # Additional Statistics
        avg_balance=avg_balance,
        active_accounts=active_accounts,
        branches=branches,

        # ==================================================
        # ML FRAUD DETECTION
        # ==================================================

        ml_analyzed=ml_analyzed,
        ml_normal=ml_normal,
        ml_suspicious=ml_suspicious,
        ml_high_risk=ml_high_risk,

        # High Risk Transactions
        high_risk_transactions=high_risk_transactions
    )

# ==========================================================
# ML ANALYTICS
# ==========================================================

@app.route("/ml-analytics")
def ml_analytics():

    connection = None
    cursor = None

    try:

        connection = get_connection()

        if not connection:
            return "Database Connection Failed"

        cursor = connection.cursor(dictionary=True)

        # ==================================================
        # MODEL INFORMATION
        # ==================================================

        model_version = "v1.0"
        model_name = "Random Forest"
        total_features = 24
        detection_threshold = 0.55

        # ==================================================
        # TRANSACTION ML STATISTICS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
        """)

        total_analyzed = cursor.fetchone()["total"] or 0

        # Normal transactions

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_prediction = 0
        """)

        normal_transactions = cursor.fetchone()["total"] or 0

        # Suspicious transactions

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_prediction = 1
        """)

        suspicious_transactions = cursor.fetchone()["total"] or 0

        # High-risk transactions

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_risk_level = 'High'
        """)

        high_risk_transactions_count = cursor.fetchone()["total"] or 0

        # Medium-risk transactions

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_risk_level = 'Medium'
        """)

        medium_risk_transactions = cursor.fetchone()["total"] or 0

        # Low-risk transactions

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_checked_at IS NOT NULL
            AND ml_risk_level = 'Low'
        """)

        low_risk_transactions = cursor.fetchone()["total"] or 0

        # ==================================================
        # SUSPICIOUS TRANSACTIONS
        # ==================================================

        cursor.execute("""
            SELECT
                t.transaction_id,
                t.transaction_reference,
                t.transaction_type,
                t.amount,
                t.transaction_date,
                t.ml_prediction,
                t.ml_risk_score,
                t.ml_risk_level,
                a.account_number

            FROM Transactions t

            LEFT JOIN Accounts a
                ON t.account_id = a.account_id

            WHERE t.ml_checked_at IS NOT NULL
            AND t.ml_prediction = 1

            ORDER BY t.ml_risk_score DESC

            LIMIT 20
        """)

        suspicious_list = cursor.fetchall()

        # ==================================================
        # HIGH-RISK TRANSACTIONS
        # ==================================================

        cursor.execute("""
            SELECT
                t.transaction_reference,
                t.transaction_type,
                t.amount,
                t.transaction_date,
                t.ml_risk_score,
                t.ml_risk_level,
                a.account_number

            FROM Transactions t

            LEFT JOIN Accounts a
                ON t.account_id = a.account_id

            WHERE t.ml_checked_at IS NOT NULL
            AND t.ml_risk_level = 'High'

            ORDER BY t.ml_risk_score DESC

            LIMIT 10
        """)

        high_risk_list = cursor.fetchall()

        # ==================================================
        # RECENT ML ANALYSIS
        # ==================================================

        cursor.execute("""
            SELECT
                t.transaction_reference,
                t.transaction_type,
                t.amount,
                t.transaction_date,
                t.ml_prediction,
                t.ml_risk_score,
                t.ml_risk_level

            FROM Transactions t

            WHERE t.ml_checked_at IS NOT NULL

            ORDER BY t.ml_checked_at DESC

            LIMIT 10
        """)

        recent_ml_transactions = cursor.fetchall()

        # ==================================================
        # MODEL PERFORMANCE
        # ==================================================

        # These values come from the trained/tested model.
        # Test-set performance previously obtained:

        accuracy = 0.980
        precision = 0.760
        recall = 0.470
        f1_score = 0.580

        return render_template(
            "ml_analytics.html",

            # Model information
            model_name=model_name,
            model_version=model_version,
            total_features=total_features,
            detection_threshold=detection_threshold,

            # ML statistics
            total_analyzed=total_analyzed,
            normal_transactions=normal_transactions,
            suspicious_transactions=suspicious_transactions,
            high_risk_transactions_count=high_risk_transactions_count,
            medium_risk_transactions=medium_risk_transactions,
            low_risk_transactions=low_risk_transactions,

            # Transaction data
            suspicious_list=suspicious_list,
            high_risk_list=high_risk_list,
            recent_ml_transactions=recent_ml_transactions,

            # Model metrics
            accuracy=accuracy,
            precision=precision,
            recall=recall,
            f1_score=f1_score
        )

    except Exception as e:

        print("ML Analytics Error:", e)

        return f"ML Analytics Error: {e}"

    finally:

        if cursor:
            cursor.close()

        if connection:
            connection.close()

# ==========================================================
# REPORTS
# ==========================================================

@app.route("/reports")
def reports():

    connection = get_connection()

    if not connection:

        return "Database Connection Failed"


    cursor = connection.cursor()


    try:

        # ==================================================
        # CUSTOMERS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Customers
        """)

        customers = cursor.fetchone()[0]


        # ==================================================
        # ACCOUNTS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Accounts
        """)

        accounts = cursor.fetchone()[0]


        # ==================================================
        # TRANSACTIONS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
        """)

        transactions = cursor.fetchone()[0]


        # ==================================================
        # EMPLOYEES
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Employees
        """)

        employees = cursor.fetchone()[0]


        # ==================================================
        # CARDS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Cards
        """)

        cards = cursor.fetchone()[0]


        # ==================================================
        # ACTIVE CARDS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Cards
            WHERE card_status = 'Active'
        """)

        active_cards = cursor.fetchone()[0]


        # ==================================================
        # TOTAL BANK BALANCE
        # ==================================================

        cursor.execute("""
            SELECT IFNULL(SUM(balance), 0)
            FROM Accounts
        """)

        total_balance = cursor.fetchone()[0]


        # ==================================================
        # APPROVED LOANS
        # ==================================================

        cursor.execute("""
            SELECT IFNULL(SUM(loan_amount), 0)
            FROM Loans
            WHERE loan_status = 'Approved'
        """)

        total_loans = cursor.fetchone()[0]


        # ==================================================
        # ML TRANSACTIONS ANALYZED
        #
        # A transaction is considered analyzed when
        # ml_risk_score is not NULL.
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
        """)

        ml_analyzed = cursor.fetchone()[0]


        # ==================================================
        # ML NORMAL TRANSACTIONS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_prediction = 0
        """)

        ml_normal = cursor.fetchone()[0]


        # ==================================================
        # ML SUSPICIOUS TRANSACTIONS
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_prediction = 1
        """)

        ml_suspicious = cursor.fetchone()[0]


        # ==================================================
        # ML LOW RISK
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'Low'
        """)

        ml_low_risk = cursor.fetchone()[0]


        # ==================================================
        # ML MEDIUM RISK
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'Medium'
        """)

        ml_medium_risk = cursor.fetchone()[0]


        # ==================================================
        # ML HIGH RISK
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'High'
        """)

        ml_high_risk = cursor.fetchone()[0]


        return render_template(

            "reports.html",

            # Banking information

            customers=customers,

            accounts=accounts,

            transactions=transactions,

            employees=employees,

            cards=cards,

            active_cards=active_cards,

            total_balance=total_balance,

            total_loans=total_loans,


            # ML information

            ml_analyzed=ml_analyzed,

            ml_normal=ml_normal,

            ml_suspicious=ml_suspicious,

            ml_low_risk=ml_low_risk,

            ml_medium_risk=ml_medium_risk,

            ml_high_risk=ml_high_risk

        )


    finally:

        cursor.close()

        connection.close()

# ==========================================================
# EXPORT BANKING REPORT TO EXCEL
# ==========================================================

@app.route("/export_excel")
def export_excel():

    connection = get_connection()

    if not connection:

        return "Database Connection Failed"


    cursor = connection.cursor(dictionary=True)


    try:

        # ==================================================
        # BANKING DATA
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Customers
        """)

        customers = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Accounts
        """)

        accounts = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
        """)

        transactions = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Employees
        """)

        employees = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Cards
        """)

        cards = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Cards
            WHERE card_status = 'Active'
        """)

        active_cards = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COALESCE(
                SUM(balance),
                0
            ) AS total
            FROM Accounts
        """)

        total_balance = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COALESCE(
                SUM(loan_amount),
                0
            ) AS total
            FROM Loans
            WHERE loan_status = 'Approved'
        """)

        total_loans = cursor.fetchone()["total"]


        # ==================================================
        # ML DATA
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
        """)

        ml_analyzed = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_prediction = 0
        """)

        ml_normal = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_prediction = 1
        """)

        ml_suspicious = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'Low'
        """)

        ml_low_risk = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'Medium'
        """)

        ml_medium_risk = cursor.fetchone()["total"]


        cursor.execute("""
            SELECT COUNT(*) AS total
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'High'
        """)

        ml_high_risk = cursor.fetchone()["total"]


    finally:

        cursor.close()

        connection.close()


    # ======================================================
    # CREATE EXCEL WORKBOOK
    # ======================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Banking Report"


    # ======================================================
    # MAIN HEADING
    # ======================================================

    worksheet["A1"] = "BANKING MANAGEMENT SYSTEM"

    worksheet["A2"] = (
        "Banking Reports & ML Fraud Detection"
    )


    # ======================================================
    # BANKING REPORT
    # ======================================================

    worksheet["A4"] = "BANKING OVERVIEW"

    worksheet["A5"] = "Report"

    worksheet["B5"] = "Value"


    banking_data = [

        ("Total Customers", customers),

        ("Total Accounts", accounts),

        ("Total Transactions", transactions),

        ("Total Employees", employees),

        ("Total Cards", cards),

        ("Active Cards", active_cards),

        ("Total Bank Balance", float(total_balance)),

        ("Approved Loan Amount", float(total_loans))

    ]


    row = 6


    for report_name, value in banking_data:

        worksheet.cell(

            row=row,

            column=1,

            value=report_name

        )

        worksheet.cell(

            row=row,

            column=2,

            value=value

        )

        row += 1


    # ======================================================
    # ML FRAUD DETECTION REPORT
    # ======================================================

    ml_start_row = row + 2


    worksheet.cell(

        row=ml_start_row,

        column=1,

        value="ML FRAUD DETECTION"

    )


    worksheet.cell(

        row=ml_start_row + 1,

        column=1,

        value="Report"

    )


    worksheet.cell(

        row=ml_start_row + 1,

        column=2,

        value="Value"

    )


    ml_data = [

        (
            "Transactions Analyzed",
            ml_analyzed
        ),

        (
            "Normal Transactions",
            ml_normal
        ),

        (
            "Suspicious Transactions",
            ml_suspicious
        ),

        (
            "Low Risk",
            ml_low_risk
        ),

        (
            "Medium Risk",
            ml_medium_risk
        ),

        (
            "High Risk",
            ml_high_risk
        ),

        (
            "Model",
            "Random Forest"
        ),

        (
            "Model Version",
            "v1.0"
        ),

        (
            "Features",
            24
        ),

        (
            "Classification Threshold",
            "55%"
        ),

        (
            "Accuracy",
            "98.0%"
        ),

        (
            "Precision",
            "76.0%"
        ),

        (
            "Recall",
            "47.0%"
        ),

        (
            "F1 Score",
            "58.0%"
        )

    ]


    row = ml_start_row + 2


    for report_name, value in ml_data:

        worksheet.cell(

            row=row,

            column=1,

            value=report_name

        )

        worksheet.cell(

            row=row,

            column=2,

            value=value

        )

        row += 1


    # ======================================================
    # FORMATTING
    # ======================================================

    worksheet.column_dimensions["A"].width = 34

    worksheet.column_dimensions["B"].width = 24


    # Main heading

    worksheet["A1"].font = worksheet["A1"].font.copy(

        bold=True,

        size=16

    )


    worksheet["A2"].font = worksheet["A2"].font.copy(

        bold=True,

        size=13

    )


    # Section headings

    worksheet["A4"].font = worksheet["A4"].font.copy(

        bold=True,

        size=13

    )


    worksheet.cell(

        row=ml_start_row,

        column=1

    ).font = worksheet.cell(

        row=ml_start_row,

        column=1

    ).font.copy(

        bold=True,

        size=13

    )


    # Table headings

    worksheet["A5"].font = worksheet["A5"].font.copy(

        bold=True

    )


    worksheet["B5"].font = worksheet["B5"].font.copy(

        bold=True

    )


    worksheet.cell(

        row=ml_start_row + 1,

        column=1

    ).font = worksheet.cell(

        row=ml_start_row + 1,

        column=1

    ).font.copy(

        bold=True

    )


    worksheet.cell(

        row=ml_start_row + 1,

        column=2

    ).font = worksheet.cell(

        row=ml_start_row + 1,

        column=2

    ).font.copy(

        bold=True

    )


    # ======================================================
    # CURRENCY FORMATTING
    # ======================================================

    # Total Bank Balance is row 12

    worksheet["B12"].number_format = '₹#,##0.00'


    # Approved Loan Amount is row 13

    worksheet["B13"].number_format = '₹#,##0.00'


    # ======================================================
    # SAVE TO MEMORY
    # ======================================================

    excel_buffer = BytesIO()


    workbook.save(excel_buffer)


    excel_buffer.seek(0)


    # ======================================================
    # DOWNLOAD
    # ======================================================

    return send_file(

        excel_buffer,

        as_attachment=True,

        download_name="Banking_ML_Report.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    )

# ==========================================================
# EXPORT PDF REPORT
# ==========================================================

@app.route("/export_pdf")
def export_pdf():

    connection = get_connection()

    if not connection:

        return "Database Connection Failed"


    cursor = connection.cursor()


    try:

        # ==================================================
        # BANKING DATA
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Customers
        """)

        customers = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Accounts
        """)

        accounts = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
        """)

        transactions = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Employees
        """)

        employees = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Cards
        """)

        cards = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Cards
            WHERE card_status = 'Active'
        """)

        active_cards = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COALESCE(
                SUM(balance),
                0
            )
            FROM Accounts
        """)

        total_balance = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COALESCE(
                SUM(loan_amount),
                0
            )
            FROM Loans
            WHERE loan_status = 'Approved'
        """)

        total_loans = cursor.fetchone()[0]


        # ==================================================
        # ML DATA
        # ==================================================

        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
        """)

        ml_analyzed = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_prediction = 0
        """)

        ml_normal = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_prediction = 1
        """)

        ml_suspicious = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'Low'
        """)

        ml_low_risk = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'Medium'
        """)

        ml_medium_risk = cursor.fetchone()[0]


        cursor.execute("""
            SELECT COUNT(*)
            FROM Transactions
            WHERE ml_risk_score IS NOT NULL
            AND ml_risk_level = 'High'
        """)

        ml_high_risk = cursor.fetchone()[0]


        # ==================================================
        # ACCOUNT DETAILS
        # ==================================================

        cursor.execute("""
            SELECT
                account_number,
                account_type,
                balance
            FROM Accounts
        """)

        account_rows = cursor.fetchall()


    finally:

        cursor.close()

        connection.close()


    # ======================================================
    # CREATE TEMPORARY PDF
    # ======================================================

    temp = tempfile.NamedTemporaryFile(

        delete=False,

        suffix=".pdf"

    )


    document = SimpleDocTemplate(

        temp.name,

        pagesize=letter,

        rightMargin=35,

        leftMargin=35,

        topMargin=35,

        bottomMargin=35

    )


    # ======================================================
    # PDF CONTENT
    # ======================================================

    pdf_elements = []


    # ======================================================
    # TITLE
    # ======================================================

    pdf_elements.append(

        Paragraph(

            "BANKING MANAGEMENT SYSTEM",

            ParagraphStyle(

                "Title",

                fontSize=18,

                leading=22,

                alignment=1,

                spaceAfter=8

            )

        )

    )


    pdf_elements.append(

        Paragraph(

            "Banking Reports & ML Fraud Detection",

            ParagraphStyle(

                "Subtitle",

                fontSize=12,

                leading=16,

                alignment=1,

                spaceAfter=20

            )

        )

    )


    # ======================================================
    # BANKING OVERVIEW
    # ======================================================

    pdf_elements.append(

        Paragraph(

            "Banking Overview",

            ParagraphStyle(

                "Heading",

                fontSize=14,

                leading=18,

                spaceAfter=10

            )

        )

    )


    banking_pdf_data = [

        ["Report", "Value"],

        ["Total Customers", customers],

        ["Total Accounts", accounts],

        ["Total Transactions", transactions],

        ["Total Employees", employees],

        ["Total Cards", cards],

        ["Active Cards", active_cards],

        [
            "Total Bank Balance",
            f"₹ {float(total_balance):,.2f}"
        ],

        [
            "Approved Loan Amount",
            f"₹ {float(total_loans):,.2f}"
        ]

    ]


    banking_table = Table(

        banking_pdf_data,

        colWidths=[300, 180]

    )


    banking_table.setStyle(

        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1e3a8a")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "ALIGN",
                (1, 1),
                (1, -1),
                "RIGHT"
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, 0),
                8
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, 0),
                8
            )

        ])

    )


    pdf_elements.append(

        banking_table

    )


    pdf_elements.append(

        Spacer(1, 25)

    )


    # ======================================================
    # ML FRAUD DETECTION
    # ======================================================

    pdf_elements.append(

        Paragraph(

            "ML Fraud Detection",

            ParagraphStyle(

                "MLHeading",

                fontSize=14,

                leading=18,

                spaceAfter=10

            )

        )

    )


    ml_pdf_data = [

        ["ML Report", "Value"],

        [
            "Transactions Analyzed",
            ml_analyzed
        ],

        [
            "Normal Transactions",
            ml_normal
        ],

        [
            "Suspicious Transactions",
            ml_suspicious
        ],

        [
            "Low Risk",
            ml_low_risk
        ],

        [
            "Medium Risk",
            ml_medium_risk
        ],

        [
            "High Risk",
            ml_high_risk
        ],

        [
            "Model",
            "Random Forest"
        ],

        [
            "Model Version",
            "v1.0"
        ],

        [
            "Features",
            "24"
        ],

        [
            "Classification Threshold",
            "55%"
        ],

        [
            "Accuracy",
            "98.0%"
        ],

        [
            "Precision",
            "76.0%"
        ],

        [
            "Recall",
            "47.0%"
        ],

        [
            "F1 Score",
            "58.0%"
        ]

    ]


    ml_table = Table(

        ml_pdf_data,

        colWidths=[300, 180]

    )


    ml_table.setStyle(

        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#374151")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "ALIGN",
                (1, 1),
                (1, -1),
                "RIGHT"
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, 0),
                8
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, 0),
                8
            )

        ])

    )


    pdf_elements.append(

        ml_table

    )


    pdf_elements.append(

        Spacer(1, 25)

    )


    # ======================================================
    # ACCOUNT DETAILS
    # ======================================================

    pdf_elements.append(

        Paragraph(

            "Account Details",

            ParagraphStyle(

                "AccountHeading",

                fontSize=14,

                leading=18,

                spaceAfter=10

            )

        )

    )


    account_pdf_data = [

        [
            "Account Number",
            "Account Type",
            "Balance"
        ]

    ]


    for row in account_rows:

        account_pdf_data.append([

            row[0],

            row[1],

            f"₹ {float(row[2]):,.2f}"

        ])


    account_table = Table(

        account_pdf_data,

        colWidths=[180, 150, 150]

    )


    account_table.setStyle(

        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1e3a8a")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "ALIGN",
                (2, 1),
                (2, -1),
                "RIGHT"
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, 0),
                8
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, 0),
                8
            )

        ])

    )


    pdf_elements.append(

        account_table

    )


    # ======================================================
    # BUILD PDF
    # ======================================================

    document.build(

        pdf_elements

    )


    # ======================================================
    # SEND PDF
    # ======================================================

    return send_file(

        temp.name,

        as_attachment=True,

        download_name="Banking_ML_Report.pdf",

        mimetype="application/pdf"

    )

# ==========================================================
# SHOW REGISTERED ROUTES
# ==========================================================

@app.route("/routes")
def show_routes():

    routes = []


    for rule in app.url_map.iter_rules():

        routes.append(str(rule))


    return "<br>".join(
        sorted(routes)
    )


# ==========================================================
# DATABASE BACKUP
# ==========================================================

@app.route("/backup_database")
def backup_database():

    backup_folder = "backup"


    # ======================================================
    # CREATE BACKUP FOLDER
    # ======================================================

    if not os.path.exists(backup_folder):

        os.makedirs(backup_folder)


    # ======================================================
    # BACKUP FILE
    # ======================================================

    backup_file = os.path.join(
        backup_folder,
        "bankingdb_backup.sql"
    )


    # ======================================================
    # MYSQL DUMP COMMAND
    # ======================================================

    command = [

        "mysqldump",

        "-u",

        "root",

        "-p" + os.environ.get(
            "MYSQL_PASSWORD",
            ""
        ),

        "bankingdb"

    ]


    # ======================================================
    # CREATE DATABASE BACKUP
    # ======================================================

    with open(
        backup_file,
        "w",
        encoding="utf-8"
    ) as outfile:

        result = subprocess.run(

            command,

            stdout=outfile,

            stderr=subprocess.PIPE,

            text=True

        )


    # ======================================================
    # BACKUP SUCCESSFUL
    # ======================================================

    if result.returncode == 0:

        print(
            "Database Backup Created Successfully."
        )


        # ==================================================
        # GET LOGGED-IN ADMIN
        # ==================================================

        username = session.get(
            "username",
            "Unknown Admin"
        )


        # ==================================================
        # ADD ACTIVITY LOG
        # ==================================================

        connection = get_connection()


        if connection:

            cursor = connection.cursor()


            cursor.execute(
                """
                INSERT INTO Activity_Log
                (
                    activity
                )
                VALUES
                (
                    %s
                )
                """,
                (
                    f"Database Backup Created by Admin '{username}'",
                )
            )


            connection.commit()

            cursor.close()

            connection.close()


    # ======================================================
    # BACKUP FAILED
    # ======================================================

    else:

        print("Backup Failed")

        print(result.stderr)


    return redirect("/dashboard")


# ==========================================================
# DOWNLOAD DATABASE BACKUP
# ==========================================================

@app.route("/download_backup")
def download_backup():

    backup_file = os.path.join(
        "backup",
        "bankingdb_backup.sql"
    )


    if os.path.exists(backup_file):

        return send_file(

            backup_file,

            as_attachment=True

        )


    return "Backup file not found."


# ==========================================================
# RUN FLASK APPLICATION
# ==========================================================

if __name__ == "__main__":

    print(
        "\n========== REGISTERED ROUTES =========="
    )


    for rule in app.url_map.iter_rules():

        print(rule)


    print(
        "=======================================\n"
    )


    app.run(
        debug=True
    )